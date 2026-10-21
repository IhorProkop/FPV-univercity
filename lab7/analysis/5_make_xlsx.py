from pathlib import Path
HERE = Path(__file__).resolve().parent          # lab7/analysis
FLIGHT = HERE.parent / "flight"                 # lab7/flight (журнали польоту)
import os; os.chdir(HERE)
import pandas as pd, json
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font, PatternFill, Alignment
df=pd.read_csv(FLIGHT / "2026-10-20 21-23-33 vehicle1.csv")
s=json.load(open('csvstats.json')); n=s['dis_row']+11
cols=['Timestamp','flightTime','missionItemIndex','altitudeRelative','airSpeed','groundSpeed','climbRate','roll','pitch','heading','gps.lat','gps.lon','distanceToHome']
d=df[cols].iloc[:n].copy(); d['Timestamp']=pd.to_datetime(d.Timestamp)
wb=Workbook(); ws=wb.active; ws.title='CSV_QGC'
hdr=cols+['t, с','Висота, м','Повітряна шв., м/с','Наземна шв., м/с']
ws.append(hdr)
for i,r in enumerate(d.itertuples(index=False),start=2):
    ws.append(list(r)+[f'=(A{i}-$A$2)*86400',f'=D{i}*0.3048',f'=E{i}*0.44704',f'=F{i}*0.44704'])
    ws[f'A{i}'].number_format='hh:mm:ss.000'
last=n+1
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='E7E6E6'); c.alignment=Alignment(wrap_text=True)
for col,w in zip('ABCDEFGHIJKLMNOPQ',[14,9,8,10,9,9,9,7,7,8,12,12,10,8,10,12,12]): ws.column_dimensions[col].width=w
ws.freeze_panes='A2'
k=wb.create_sheet('Показники',0)
R=f'$2:$'  # unused
rows=[('Показник','Значення','Формула / пояснення'),
('Час армування (початок запису CSV)','=CSV_QGC!A2','перший рядок CSV'),
('Момент відриву від землі (висота > 0,3 м)',f'=_xlfn.MINIFS(CSV_QGC!A2:A{last},CSV_QGC!O2:O{last},">0.3")','MINIFS'),
('Час торкання землі',f'=_xlfn.MAXIFS(CSV_QGC!A2:A{last},CSV_QGC!O2:O{last},">0.3")','останній рядок з висотою > 0,3 м'),
('Найбільша висота, м',f'=MAX(CSV_QGC!O2:O{last})','altitudeRelative × 0,3048'),
('Найбільша повітряна швидкість, м/с',f'=MAX(CSV_QGC!P2:P{last})','airSpeed × 0,44704'),
('Найбільша наземна швидкість, м/с',f'=MAX(CSV_QGC!Q2:Q{last})','groundSpeed × 0,44704'),
('Від arm до disarm за годинником ПК, с',f'=(CSV_QGC!A{s["dis_row"]+2}-CSV_QGC!A2)*86400','рядок, де flightTime перестає зростати'),
('Від зльоту до посадки за годинником ПК, с','=(B4-B3)*86400','реальний час, RTF ≈ 0,2'),
('Початкова координата lat',f'=CSV_QGC!K2',''),('Початкова координата lon',f'=CSV_QGC!L2',''),
('Кінцева координата lat',f'=CSV_QGC!K{s["dis_row"]+2}',''),('Кінцева координата lon',f'=CSV_QGC!L{s["dis_row"]+2}',''),
('Максимальний |roll|, град.',f'=MAX(MAX(CSV_QGC!H2:H{last}),-MIN(CSV_QGC!H2:H{last}))',''),
('Середня висота на ділянці WP, м',f'=AVERAGEIF(CSV_QGC!C2:C{last},2,CSV_QGC!O2:O{last})','missionItemIndex = 2'),
('Відстань від H після посадки, м',f'=CSV_QGC!M{s["dis_row"]+2}*0.3048','distanceToHome × 0,3048'),
('Час зльоту (час автопілота), с',20.6,'з .tlog: набір 0,2 → 44,9 м без зависання на 3 м'),('Час польоту (час автопілота), с',131.8,'з .tlog: від відриву до торкання'),('Режими польоту','GUIDED → AUTO','HEARTBEAT.custom_mode: 4 → 3 (.tlog)'),('Стан arm до / після місії','DISARMED / DISARMED','HEARTBEAT.base_mode: 81 / 89, біт 128 = 0')]
for r in rows: k.append(r)
for c in k[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='E7E6E6')
for a in ('B2','B3','B4'): k[a].number_format='hh:mm:ss'
for a in ('B5','B6','B7','B8','B9','B14','B15','B16'): k[a].number_format='0.00'
for a in ('B17','B18'): k[a].number_format='0.0'; k[a].font=Font(color='0000FF')
for a in ('B10','B11','B12','B13'): k[a].number_format='0.0000000'
k.column_dimensions['A'].width=38; k.column_dimensions['B'].width=30; k.column_dimensions['C'].width=40
g=wb.create_sheet('Графіки',1)
x=Reference(ws,min_col=14,min_row=2,max_row=last)
for i,(col,title,yt) in enumerate([(15,'Висота від часу','Висота, м'),(16,'Повітряна швидкість від часу','м/с'),(17,'Наземна швидкість від часу','м/с')]):
    ch=ScatterChart(); ch.title=title; ch.style=13; ch.x_axis.title='Час від армування, с'; ch.y_axis.title=yt
    ser=Series(Reference(ws,min_col=col,min_row=2,max_row=last),x,title=title); ser.marker.symbol='none'; ser.smooth=False
    ch.series.append(ser); ch.legend=None; ch.width=24; ch.height=8.5
    ch.x_axis.delete=False; ch.y_axis.delete=False
    g.add_chart(ch,f'A{1+i*18}')
wb.save('Аналіз_телеметрії_ЛР7.xlsx')
